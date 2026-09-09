"""S5-T2 RED: Document and table parser contract tests.

Verifies the document→section→paragraph/table→row/cell/code-block hierarchy,
stable locators, deterministic content digests, and locator replay.
"""

from __future__ import annotations

from pathlib import Path

from zhiwei.knowledge.connectors.files import FileConnector
from zhiwei.knowledge.parsers.documents import DocumentParser, NodeType
from zhiwei.knowledge.parsers.tables import TableParser

FIXTURES = Path(__file__).resolve().parent.parent.parent / "fixtures" / "knowledge"


# ---------------------------------------------------------------------------
# Document parser — Markdown
# ---------------------------------------------------------------------------

class TestMarkdownDocumentParser:
    def test_parse_sample_md(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc = parser.parse_file(path)
        assert doc.format == "markdown"
        assert doc.title == "Sample Document"

    def test_root_is_document_node(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc = parser.parse_file(path)
        assert doc.root.node_type == NodeType.DOCUMENT

    def test_sections_extracted(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc = parser.parse_file(path)
        all_nodes = doc.flat_nodes()
        sections = [n for n in all_nodes if n.node_type == NodeType.SECTION]
        titles = [s.title for s in sections]
        assert "Section One" in titles
        assert "Section Two" in titles

    def test_code_block_extracted(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc = parser.parse_file(path)
        all_nodes = doc.flat_nodes()
        code_blocks = [n for n in all_nodes if n.node_type == NodeType.CODE_BLOCK]
        assert len(code_blocks) == 1
        assert "def hello" in code_blocks[0].content
        assert code_blocks[0].title == "python"

    def test_paragraphs_extracted(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc = parser.parse_file(path)
        all_nodes = doc.flat_nodes()
        paras = [n for n in all_nodes if n.node_type == NodeType.PARAGRAPH]
        assert len(paras) >= 3

    def test_content_digest_is_sha256(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc = parser.parse_file(path)
        assert doc.content_digest.startswith("sha256:")
        assert len(doc.content_digest) == 71

    def test_digest_deterministic(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc_a = parser.parse_file(path)
        doc_b = parser.parse_file(path)
        assert doc_a.content_digest == doc_b.content_digest

    def test_digest_changes_with_content(self) -> None:
        parser = DocumentParser()
        doc_a = parser.parse_markdown(b"# Title\n", source_uri="u1")
        doc_b = parser.parse_markdown(b"# Different\n", source_uri="u1")
        assert doc_a.content_digest != doc_b.content_digest

    def test_locators_are_unique(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc = parser.parse_file(path)
        uris = [n.locator.uri for n in doc.flat_nodes()]
        assert len(uris) == len(set(uris))

    def test_locator_replay(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        raw = path.read_bytes()
        parser = DocumentParser()
        doc = parser.parse_file(path)
        replayed = doc.replay_locators(raw)
        assert len(replayed) > 0
        for _uri, content in replayed.items():
            assert isinstance(content, str)

    def test_root_locator_matches_source(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc = parser.parse_file(path)
        assert doc.source_locator.uri == f"file://{path.resolve()}"
        assert doc.root.locator == doc.source_locator

    def test_span_populated_for_sections(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        parser = DocumentParser()
        doc = parser.parse_file(path)
        sections = [c for c in doc.root.children if c.node_type == NodeType.SECTION]
        for section in sections:
            assert section.span is not None
            assert section.span[0] < section.span[1]

    def test_format_markdown(self) -> None:
        parser = DocumentParser()
        doc = parser.parse_markdown(b"# Hi\n\nBody.\n", source_uri="test://x")
        assert doc.format == "markdown"


# ---------------------------------------------------------------------------
# Document parser — CSV
# ---------------------------------------------------------------------------

class TestCSVDocumentParser:
    def test_csv_to_document(self) -> None:
        csv_data = b"name,age\nAlice,30\nBob,25\n"
        parser = DocumentParser()
        doc = parser.parse_csv(csv_data, source_uri="test://csv")
        assert doc.format == "csv"
        all_nodes = doc.flat_nodes()
        tables = [n for n in all_nodes if n.node_type == NodeType.TABLE]
        assert len(tables) == 1
        assert "Alice" in tables[0].content

    def test_empty_csv(self) -> None:
        parser = DocumentParser()
        doc = parser.parse_csv(b"", source_uri="test://empty")
        assert doc.root.node_type == NodeType.DOCUMENT
        assert len(doc.root.children) == 0


# ---------------------------------------------------------------------------
# Table parser — CSV
# ---------------------------------------------------------------------------

class TestCSVTableParser:
    def test_parse_csv(self) -> None:
        csv_data = b"name,age\nAlice,30\nBob,25\n"
        parser = TableParser()
        table = parser.parse_csv(csv_data, source_uri="test://csv")
        assert table.headers == ("name", "age")
        assert len(table.rows) == 3  # header + 2 data rows

    def test_cell_locators(self) -> None:
        csv_data = b"name,age\nAlice,30\n"
        parser = TableParser()
        table = parser.parse_csv(csv_data, source_uri="test://csv")
        first_cell = table.rows[0].cells[0]
        assert first_cell.locator.uri == "test://csv#row=0&col=0"
        assert first_cell.value == "name"

    def test_content_digest_deterministic(self) -> None:
        csv_data = b"name,age\nAlice,30\n"
        parser = TableParser()
        t1 = parser.parse_csv(csv_data, source_uri="test://x")
        t2 = parser.parse_csv(csv_data, source_uri="test://x")
        assert t1.content_digest == t2.content_digest

    def test_digest_differs_for_different_data(self) -> None:
        parser = TableParser()
        t1 = parser.parse_csv(b"a,b\n1,2\n", source_uri="x")
        t2 = parser.parse_csv(b"a,b\n3,4\n", source_uri="x")
        assert t1.content_digest != t2.content_digest

    def test_locator_replay(self) -> None:
        csv_data = b"col1,col2\nval1,val2\n"
        parser = TableParser()
        table = parser.parse_csv(csv_data, source_uri="test://replay")
        replayed = table.replay_locators(csv_data)
        assert len(replayed) > 0
        vals = set(replayed.values())
        assert "col1" in vals
        assert "val2" in vals

    def test_tsv_delimiter(self) -> None:
        tsv_data = b"name\tage\nAlice\t30\n"
        parser = TableParser()
        table = parser.parse_csv(tsv_data, source_uri="test://tsv", delimiter="\t")
        assert table.headers == ("name", "age")
        assert table.rows[1].cells[0].value == "Alice"


# ---------------------------------------------------------------------------
# Table parser — XLSX
# ---------------------------------------------------------------------------

class TestXLSXTableParser:
    def test_parse_xlsx(self) -> None:
        import io

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Data"
        ws.append(["name", "age"])
        ws.append(["Alice", 30])
        ws.append(["Bob", 25])
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        data = buf.getvalue()

        parser = TableParser()
        table = parser.parse_xlsx(data, source_uri="test://xlsx")
        assert table.headers == ("name", "age")
        assert len(table.rows) == 3
        assert table.sheet_name == "Data"

    def test_xlsx_digest_deterministic(self) -> None:
        import io

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["x", "y"])
        ws.append([1, 2])
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        data = buf.getvalue()

        parser = TableParser()
        t1 = parser.parse_xlsx(data, source_uri="x")
        t2 = parser.parse_xlsx(data, source_uri="x")
        assert t1.content_digest == t2.content_digest


# ---------------------------------------------------------------------------
# File connector
# ---------------------------------------------------------------------------

class TestFileConnector:
    def test_read_and_digest(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        connector = FileConnector()
        snap = connector.read(path)
        assert snap.content_digest.startswith("sha256:")
        assert snap.size_bytes > 0
        assert snap.locator.uri == f"file://{path.resolve()}"

    def test_digest_deterministic(self) -> None:
        path = FIXTURES / "documents" / "sample.md"
        connector = FileConnector()
        s1 = connector.read(path)
        s2 = connector.read(path)
        assert s1.content_digest == s2.content_digest

    def test_to_source_version(self) -> None:
        from uuid import uuid4

        path = FIXTURES / "documents" / "sample.md"
        connector = FileConnector()
        snap = connector.read(path)
        sv = connector.to_source_version(snap, source_object_id=uuid4())
        assert sv.content_digest == snap.content_digest
        assert sv.locator == snap.locator
        assert sv.version_seq == 1

    def test_to_source_object(self) -> None:
        from uuid import uuid4

        path = FIXTURES / "documents" / "sample.md"
        connector = FileConnector()
        snap = connector.read(path)
        org_id = uuid4()
        ws_id = uuid4()
        obj = connector.to_source_object(snap, organization_id=org_id, workspace_id=ws_id)
        assert obj.organization_id == org_id
        assert obj.workspace_id == ws_id
        assert obj.source_type == "file"
